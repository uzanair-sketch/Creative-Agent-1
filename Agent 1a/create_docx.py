import json
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    """Sets background shading for a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tc_pr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets cell padding."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tc_pr.append(tc_mar)

def create_treatment_docx(json_path="output/treatment_output.json", docx_path="output/treatment.docx"):
    with open(json_path, "r") as f:
        data = json.load(f)

    doc = Document()

    # Set page margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Style definitions
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Arial'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    # Header / Title
    treatment = data.get("treatment", {})
    brief = data.get("brief_analysis", {})

    title_p = doc.add_paragraph()
    title_run = title_p.add_run(f"Commercial Treatment: {treatment.get('treatment_title', 'Untitled')}")
    title_run.font.size = Pt(22)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(0x0f, 0x17, 0x2a) # Dark slate
    title_p.paragraph_format.space_after = Pt(4)

    subtitle_p = doc.add_paragraph()
    sub_run = subtitle_p.add_run("Client Pitch & Production Treatment Document")
    sub_run.font.size = Pt(11)
    sub_run.font.italic = True
    sub_run.font.color.rgb = RGBColor(0x64, 0x74, 0x8b)
    subtitle_p.paragraph_format.space_after = Pt(16)

    # Metadata Box
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    meta_items = [
        ("Client", brief.get("client_name", "")),
        ("Objective", brief.get("core_objective", "")),
        ("Target Audience", brief.get("target_audience", "")),
        ("Aesthetic Tone", brief.get("aesthetic_tone", ""))
    ]

    for i, (label, val) in enumerate(meta_items):
        row = meta_table.rows[i]
        c0 = row.cells[0]
        c1 = row.cells[1]
        c0.width = Inches(1.8)
        c1.width = Inches(5.1)

        set_cell_background(c0, "F1F5F9")
        set_cell_background(c1, "F8FAFC")
        set_cell_margins(c0, top=80, bottom=80, left=120, right=120)
        set_cell_margins(c1, top=80, bottom=80, left=120, right=120)

        p0 = c0.paragraphs[0]
        r0 = p0.add_run(label)
        r0.bold = True
        r0.font.size = Pt(10)
        r0.font.color.rgb = RGBColor(0x1e, 0x29, 0x3b)

        p1 = c1.paragraphs[0]
        r1 = p1.add_run(val)
        r1.font.size = Pt(10)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Section 1: Creative Concept & Logline
    h1 = doc.add_heading(level=1)
    r_h1 = h1.add_run("1. Creative Concept & Logline")
    r_h1.font.size = Pt(14)
    r_h1.font.color.rgb = RGBColor(0x0f, 0x17, 0x2a)

    # Logline Callout box
    logline_tbl = doc.add_table(rows=1, cols=1)
    logline_cell = logline_tbl.rows[0].cells[0]
    logline_cell.width = Inches(6.9)
    set_cell_background(logline_cell, "F0FDF4")
    set_cell_margins(logline_cell, top=120, bottom=120, left=160, right=160)
    lp = logline_cell.paragraphs[0]
    lr1 = lp.add_run("Logline: ")
    lr1.bold = True
    lr1.font.color.rgb = RGBColor(0x16, 0x65, 0x34)
    lr2 = lp.add_run(treatment.get("logline", ""))
    lr2.italic = True
    lr2.font.color.rgb = RGBColor(0x16, 0x65, 0x34)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    p_vd_title = doc.add_paragraph()
    r_vd_title = p_vd_title.add_run("Visual Direction:")
    r_vd_title.bold = True
    p_vd_title.paragraph_format.space_after = Pt(2)

    p_vd = doc.add_paragraph(treatment.get("visual_direction", ""))
    p_vd.paragraph_format.space_after = Pt(14)

    # Section 2: Pre-Visualization Prompts
    h2 = doc.add_heading(level=1)
    r_h2 = h2.add_run("2. Pre-Visualization Prompts (Midjourney / Flux)")
    r_h2.font.size = Pt(14)
    r_h2.font.color.rgb = RGBColor(0x0f, 0x17, 0x2a)

    for prompt_text in treatment.get("ai_previs_prompts", []):
        p_pr = doc.add_paragraph(style='List Bullet')
        r_pr = p_pr.add_run(prompt_text)
        r_pr.font.size = Pt(9.5)
        r_pr.font.name = "Consolas"
        r_pr.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
        p_pr.paragraph_format.space_after = Pt(4)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Section 3: Shot List
    h3 = doc.add_heading(level=1)
    r_h3 = h3.add_run("3. Shot List")
    r_h3.font.size = Pt(14)
    r_h3.font.color.rgb = RGBColor(0x0f, 0x17, 0x2a)

    shots = data.get("shot_list", [])
    shot_table = doc.add_table(rows=len(shots) + 1, cols=6)
    shot_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    headers = ["Shot #", "Location", "Framing & Lens", "Movement", "Action", "Audio / VO"]
    widths = [Inches(0.6), Inches(1.2), Inches(1.2), Inches(1.0), Inches(1.6), Inches(1.3)]

    # Header row
    hdr_row = shot_table.rows[0]
    for j, (hdr_text, w) in enumerate(zip(headers, widths)):
        cell = hdr_row.cells[j]
        cell.width = w
        set_cell_background(cell, "1E293B")
        set_cell_margins(cell, top=80, bottom=80, left=80, right=80)
        p = cell.paragraphs[0]
        r = p.add_run(hdr_text)
        r.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0xff, 0xff, 0xff)

    for i, s in enumerate(shots):
        row = shot_table.rows[i + 1]
        bg = "F8FAFC" if i % 2 == 1 else "FFFFFF"
        vals = [
            str(s.get("shot_number", i + 1)),
            s.get("scene_location", ""),
            s.get("framing_and_lens", ""),
            s.get("camera_movement", ""),
            s.get("visual_action", ""),
            s.get("audio_vo", "")
        ]
        for j, (val, w) in enumerate(zip(vals, widths)):
            cell = row.cells[j]
            cell.width = w
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(8.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    # Section 4: Production Budget Estimate
    h4 = doc.add_heading(level=1)
    r_h4 = h4.add_run("4. Production Budget Estimate")
    r_h4.font.size = Pt(14)
    r_h4.font.color.rgb = RGBColor(0x0f, 0x17, 0x2a)

    budgets = data.get("budget_breakdown", [])
    b_table = doc.add_table(rows=len(budgets) + 2, cols=3)
    b_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    b_headers = ["Category", "Line Item Description", "Estimated Cost (USD)"]
    b_widths = [Inches(1.8), Inches(3.6), Inches(1.5)]

    # Header
    for j, (h_text, w) in enumerate(zip(b_headers, b_widths)):
        cell = b_table.rows[0].cells[j]
        cell.width = w
        set_cell_background(cell, "1E293B")
        set_cell_margins(cell, top=80, bottom=80, left=80, right=80)
        p = cell.paragraphs[0]
        r = p.add_run(h_text)
        r.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0xff, 0xff, 0xff)

    for i, b in enumerate(budgets):
        row = b_table.rows[i + 1]
        bg = "F8FAFC" if i % 2 == 1 else "FFFFFF"
        cost = f"${b.get('estimated_cost_usd', 0):,.2f}"
        row_vals = [b.get("category", ""), b.get("line_item", ""), cost]
        for j, (val, w) in enumerate(zip(row_vals, b_widths)):
            cell = row.cells[j]
            cell.width = w
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
            p = cell.paragraphs[0]
            if j == 2:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            r = p.add_run(val)
            r.font.size = Pt(9)

    # Total Row
    tot_row = b_table.rows[-1]
    tot_cost = f"${data.get('total_estimated_budget_usd', 0):,.2f}"
    set_cell_background(tot_row.cells[0], "E2E8F0")
    set_cell_background(tot_row.cells[1], "E2E8F0")
    set_cell_background(tot_row.cells[2], "E2E8F0")
    for cell in tot_row.cells:
        set_cell_margins(cell, top=80, bottom=80, left=80, right=80)

    p_tot_label = tot_row.cells[1].paragraphs[0]
    r_tot_label = p_tot_label.add_run("Total Estimated Production Budget:")
    r_tot_label.bold = True
    r_tot_label.font.size = Pt(9.5)

    p_tot_val = tot_row.cells[2].paragraphs[0]
    p_tot_val.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_tot_val = p_tot_val.add_run(tot_cost)
    r_tot_val.bold = True
    r_tot_val.font.size = Pt(9.5)
    r_tot_val.font.color.rgb = RGBColor(0x0f, 0x76, 0x6e) # Teal

    doc.save(docx_path)
    print(f"✅ Generated DOCX for Google Docs: {docx_path}")

if __name__ == "__main__":
    create_treatment_docx()
